"""Lê as planilhas como vieram: .xlsx (primeira aba) e .csv."""

import csv
import io
import zipfile
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

EXTENSIONS = {'.xlsx', '.xlsm', '.csv'}


class ReadError(Exception):
    pass


def list_spreadsheets(folder: Path) -> list[Path]:
    files = [
        path for path in folder.iterdir()
        if path.is_file() and path.suffix.lower() in EXTENSIONS
        # "~$arquivo.xlsx" é o arquivo temporário que o Excel cria enquanto
        # a planilha está aberta.
        and not path.name.startswith('~$')
    ]
    return sorted(files, key=lambda path: path.name.lower())


def read_rows(path: Path) -> list[list]:
    if path.suffix.lower() == '.csv':
        return _read_csv(path)
    return _read_xlsx(path)


def _read_xlsx(path: Path) -> list[list]:
    try:
        # data_only: o valor calculado das fórmulas, não a fórmula.
        workbook = load_workbook(path, read_only=True, data_only=True)
    except (zipfile.BadZipFile, InvalidFileException, KeyError):
        raise ReadError(
            'não consegui abrir o arquivo: ele está corrompido ou não é uma '
            'planilha do Excel',
        ) from None
    try:
        sheet = workbook.worksheets[0]
        return [list(row) for row in sheet.iter_rows(values_only=True)]
    finally:
        workbook.close()


def _read_csv(path: Path) -> list[list]:
    data = path.read_bytes()
    # O Excel em português salva em cp1252; o Google Planilhas e a maioria
    # dos sistemas, em UTF-8.
    try:
        text = data.decode('utf-8-sig')
    except UnicodeDecodeError:
        text = data.decode('cp1252')
    if not text.strip():
        return []
    try:
        delimiter = csv.Sniffer().sniff(
            text[:4096], delimiters=';,\t',
        ).delimiter
    except csv.Error:
        delimiter = ';'
    return [list(row) for row in csv.reader(io.StringIO(text),
                                            delimiter=delimiter)]
