"""Monta o relatório em Excel a partir das vendas consolidadas."""

from dataclasses import asdict
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from sales_report.cleaning import Problem, Sale

# O Excel em português mostra R$ 1.234,56 com este formato.
CURRENCY = '"R$" #,##0.00'
PERCENT = '0.0%'
DATE = 'dd/mm/yyyy'
MONTH = 'mm/yyyy'

HEADER_FONT = Font(bold=True, color='FFFFFF')
HEADER_FILL = PatternFill('solid', fgColor='4F46E5')
TITLE_FONT = Font(bold=True, size=14)


def write_report(
    path: Path,
    sales: list[Sale],
    problems: list[Problem],
    files: list[str],
) -> None:
    data = pd.DataFrame([asdict(sale) for sale in sales])
    revenue = data['total'].sum()

    workbook = Workbook()
    _summary(workbook.active, data, revenue, problems, files)
    _products(workbook.create_sheet('Produtos'), data, revenue)
    _months(workbook.create_sheet('Por mês'), data)
    _sales(workbook.create_sheet('Vendas'), sales)
    _problems(workbook.create_sheet('Problemas'), problems)
    workbook.save(path)


def _summary(sheet, data, revenue, problems, files) -> None:
    sheet.title = 'Resumo'
    sheet['A1'] = 'Relatório de vendas'
    sheet['A1'].font = TITLE_FONT

    start, end = data['date'].min(), data['date'].max()
    # Problema sem linha é a planilha inteira recusada.
    ignored = sum(1 for problem in problems if problem.line is None)
    indicators = [
        ('Período', f'{start:%d/%m/%Y} a {end:%d/%m/%Y}', None),
        ('Planilhas lidas', len(files) - ignored, None),
        ('Planilhas ignoradas', ignored, None),
        ('Vendas', len(data), None),
        ('Itens vendidos', _number(data['quantity'].sum()), None),
        ('Faturamento', revenue, CURRENCY),
        ('Ticket médio', revenue / len(data), CURRENCY),
        ('Linhas com problema', len(problems) - ignored, None),
    ]
    for row, (label, value, number_format) in enumerate(indicators, start=3):
        sheet.cell(row, 1, label).font = Font(bold=True)
        cell = sheet.cell(row, 2, value)
        cell.alignment = Alignment(horizontal='right')
        if number_format:
            cell.number_format = number_format

    by_store = (
        data.groupby('store')
        .agg(sales=('total', 'size'), items=('quantity', 'sum'),
             revenue=('total', 'sum'))
        .reset_index()
        .sort_values(['revenue', 'store'], ascending=[False, True])
    )
    first = len(indicators) + 4
    _table(sheet, first, ['Loja', 'Vendas', 'Itens', 'Faturamento',
                          '% do faturamento'], [
        [row.store, row.sales, _number(row.items), row.revenue,
         row.revenue / revenue]
        for row in by_store.itertuples()
    ], formats={4: CURRENCY, 5: PERCENT})

    last = first + len(by_store)
    chart = BarChart()
    chart.title = 'Faturamento por loja'
    chart.legend = None
    chart.add_data(Reference(sheet, min_col=4, min_row=first, max_row=last),
                   titles_from_data=True)
    chart.set_categories(Reference(sheet, min_col=1, min_row=first + 1,
                                   max_row=last))
    sheet.add_chart(chart, 'G3')
    _fit_columns(sheet)


def _products(sheet, data, revenue) -> None:
    by_product = (
        data.groupby('product')
        .agg(quantity=('quantity', 'sum'), revenue=('total', 'sum'))
        .reset_index()
        .sort_values(['revenue', 'product'], ascending=[False, True])
    )
    _table(sheet, 1, ['Produto', 'Quantidade', 'Faturamento',
                      '% do faturamento'], [
        [row.product, _number(row.quantity), row.revenue,
         row.revenue / revenue]
        for row in by_product.itertuples()
    ], formats={3: CURRENCY, 4: PERCENT})
    _fit_columns(sheet)


def _months(sheet, data) -> None:
    months = data.assign(month=data['date'].map(lambda day: day.replace(
        day=1,
    )))
    by_month = (
        months.groupby('month')
        .agg(sales=('total', 'size'), items=('quantity', 'sum'),
             revenue=('total', 'sum'))
        .reset_index()
        .sort_values('month')
    )
    _table(sheet, 1, ['Mês', 'Vendas', 'Itens', 'Faturamento'], [
        [row.month, row.sales, _number(row.items), row.revenue]
        for row in by_month.itertuples()
    ], formats={1: MONTH, 4: CURRENCY})

    chart = BarChart()
    chart.title = 'Faturamento por mês'
    chart.legend = None
    last = 1 + len(by_month)
    chart.add_data(Reference(sheet, min_col=4, min_row=1, max_row=last),
                   titles_from_data=True)
    chart.set_categories(Reference(sheet, min_col=1, min_row=2,
                                   max_row=last))
    sheet.add_chart(chart, 'F2')
    _fit_columns(sheet)


def _sales(sheet, sales) -> None:
    _table(sheet, 1, ['Data', 'Loja', 'Produto', 'Quantidade',
                      'Preço unitário', 'Total', 'Planilha', 'Linha'], [
        [sale.date, sale.store, sale.product, _number(sale.quantity),
         sale.unit_price, sale.total, sale.source, sale.line]
        for sale in sorted(sales, key=lambda sale: (sale.date, sale.store))
    ], formats={1: DATE, 5: CURRENCY, 6: CURRENCY})
    _fit_columns(sheet)


def _problems(sheet, problems) -> None:
    rows = [[problem.source, problem.line, problem.message]
            for problem in problems]
    _table(sheet, 1, ['Planilha', 'Linha', 'Problema'],
           rows or [['Nenhum problema encontrado.', None, None]])
    _fit_columns(sheet)


def _table(sheet: Worksheet, first_row: int, headers: list[str],
           rows: list[list], formats: dict[int, str] | None = None) -> None:
    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(first_row, column, header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
    for row_number, row in enumerate(rows, start=first_row + 1):
        for column, value in enumerate(row, start=1):
            cell = sheet.cell(row_number, column, value)
            if formats and column in formats:
                cell.number_format = formats[column]
    if first_row == 1:
        sheet.freeze_panes = 'A2'


def _number(value: float) -> int | float:
    # 3 itens, não 3,0.
    return int(value) if float(value).is_integer() else float(value)


def _fit_columns(sheet: Worksheet) -> None:
    widths: dict[int, int] = {}
    for row in sheet.iter_rows():
        for cell in row:
            if cell.value is not None:
                widths[cell.column] = max(widths.get(cell.column, 0),
                                          len(str(cell.value)))
    for column, width in widths.items():
        sheet.column_dimensions[get_column_letter(column)].width = min(
            max(width + 2, 10), 60,
        )
