import runpy
import sys

import pytest
from openpyxl import load_workbook

from sales_report import cli
from tests.helpers import write_csv, write_xlsx

HEADER = ['Data', 'Produto', 'Qtd', 'Preço']


@pytest.fixture
def folder(tmp_path):
    folder = tmp_path / 'vendas'
    folder.mkdir()
    write_xlsx(folder / 'loja_centro.xlsx', [
        HEADER,
        ['27/10/2026', 'Caneca', 2, 30],
        ['28/10/2026', 'Prato', 1, '1.000,00'],
    ])
    write_csv(folder / 'loja_norte.csv', [
        HEADER,
        ['27/10/2026', 'caneca', '1', '30,00'],
        ['ontem', 'Prato', '1', '10'],
    ])
    return folder


def test_generates_the_report(folder, tmp_path, capsys):
    output = tmp_path / 'saida.xlsx'

    code = cli.main([str(folder), '-o', str(output)])

    out = capsys.readouterr().out
    assert code == 0
    assert load_workbook(output).sheetnames[0] == 'Resumo'
    assert f'Lendo 2 planilhas em {folder}' in out
    assert '  loja_centro.xlsx: 2 vendas' in out
    assert '  loja_norte.csv: 1 venda, 1 linha com problema' in out
    assert f'Relatório salvo em {output}' in out
    assert '3 vendas, R$ 1.090,00 de faturamento.' in out
    assert '1 problema listado na aba "Problemas".' in out


def test_report_goes_to_the_current_folder_by_default(
    folder, tmp_path, monkeypatch,
):
    monkeypatch.chdir(tmp_path)

    assert cli.main([str(folder)]) == 0
    assert (tmp_path / 'relatorio.xlsx').exists()


def test_ignored_file_is_explained(folder, tmp_path, capsys):
    write_xlsx(folder / 'estoque.xlsx', [['Produto', 'Saldo']])

    cli.main([str(folder), '-o', str(tmp_path / 'r.xlsx')])

    out = capsys.readouterr().out
    assert ('  estoque.xlsx: ignorada (faltam as colunas: Data, Quantidade, '
            'Preço unitário ou Total)') in out


def test_missing_folder(tmp_path, capsys):
    code = cli.main([str(tmp_path / 'nao-existe')])

    assert code == 1
    assert 'Pasta não encontrada' in capsys.readouterr().err


def test_folder_without_spreadsheets(tmp_path, capsys):
    (tmp_path / 'notas.txt').write_text('x')

    code = cli.main([str(tmp_path)])

    assert code == 1
    assert 'Nenhuma planilha (.xlsx ou .csv) em' in capsys.readouterr().err


def test_no_valid_sales_writes_nothing(tmp_path, capsys):
    write_xlsx(tmp_path / 'ruim.xlsx', [HEADER, ['ontem', '', 0, 1]])
    output = tmp_path / 'r.xlsx'

    code = cli.main([str(tmp_path), '-o', str(output)])

    assert code == 1
    assert 'Nenhuma venda válida' in capsys.readouterr().err
    assert not output.exists()


def test_report_open_in_excel(folder, tmp_path, monkeypatch, capsys):
    # No Windows, o Excel trava o arquivo aberto e salvar dá
    # PermissionError.
    def locked(*args, **kwargs):
        raise PermissionError

    monkeypatch.setattr(cli, 'write_report', locked)

    code = cli.main([str(folder), '-o', str(tmp_path / 'r.xlsx')])

    assert code == 1
    assert 'Ele está aberto no Excel?' in capsys.readouterr().err


def test_runs_with_python_dash_m(folder, tmp_path, monkeypatch):
    output = tmp_path / 'r.xlsx'
    monkeypatch.setattr(sys, 'argv', ['sales_report', str(folder), '-o',
                                      str(output)])

    with pytest.raises(SystemExit) as exit_info:
        runpy.run_module('sales_report', run_name='__main__')

    assert exit_info.value.code == 0
    assert output.exists()


def test_help_is_in_portuguese(capsys):
    with pytest.raises(SystemExit):
        cli.main(['--help'])

    assert 'Junta as planilhas de vendas' in capsys.readouterr().out


def test_clean_folder_has_no_problem_line(tmp_path, capsys):
    write_xlsx(tmp_path / 'loja.xlsx', [HEADER, ['27/10/2026', 'Caneca', 1,
                                                 10]])

    cli.main([str(tmp_path), '-o', str(tmp_path / 'r.xlsx')])

    out = capsys.readouterr().out
    assert '1 venda, R$ 10,00 de faturamento.' in out
    assert 'Problemas' not in out
