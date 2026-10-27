from datetime import date

import pytest
from openpyxl import load_workbook

from sales_report.cleaning import Problem, Sale
from sales_report.report import CURRENCY, write_report

SALES = [
    Sale('centro.xlsx', 2, 'Loja Centro', date(2026, 9, 30), 'Caneca',
         2, 30.0, 60.0),
    Sale('centro.xlsx', 3, 'Loja Centro', date(2026, 10, 1), 'Camiseta',
         1, 50.0, 50.0),
    Sale('norte.csv', 2, 'Loja Norte', date(2026, 10, 15), 'Caneca',
         5, 30.0, 150.0),
    Sale('norte.csv', 3, 'Loja Norte', date(2026, 10, 31), 'Prato',
         4, 10.0, 40.0),
]
PROBLEMS = [
    Problem('norte.csv', 4, 'Produto vazio'),
    Problem('quebrada.xlsx', None, 'arquivo ignorado: a planilha está vazia'),
]
FILES = ['centro.xlsx', 'norte.csv', 'quebrada.xlsx']


@pytest.fixture
def workbook(tmp_path):
    path = tmp_path / 'relatorio.xlsx'
    write_report(path, SALES, PROBLEMS, FILES)
    return load_workbook(path)


def values(sheet, first_row=1):
    return [list(row) for row in sheet.iter_rows(min_row=first_row,
                                                 values_only=True)]


def test_sheets(workbook):
    assert workbook.sheetnames == [
        'Resumo', 'Produtos', 'Por mês', 'Vendas', 'Problemas',
    ]


def test_summary_numbers(workbook):
    sheet = workbook['Resumo']
    summary = {row[0]: row[1] for row in values(sheet) if row[0]}

    assert summary['Período'] == '30/09/2026 a 31/10/2026'
    # A planilha recusada não conta como lida nem como linha com problema.
    assert summary['Planilhas lidas'] == 2
    assert summary['Planilhas ignoradas'] == 1
    assert summary['Vendas'] == 4
    assert summary['Itens vendidos'] == 12
    assert summary['Faturamento'] == 300.0
    assert summary['Ticket médio'] == 75.0
    assert summary['Linhas com problema'] == 1


def test_money_is_formatted_in_reais(workbook):
    sheet = workbook['Resumo']
    [cell] = [row[1] for row in sheet.iter_rows()
              if row[0].value == 'Faturamento']

    assert cell.number_format == CURRENCY


def test_summary_has_revenue_by_store_with_a_chart(workbook):
    sheet = workbook['Resumo']
    rows = values(sheet)
    start = next(i for i, row in enumerate(rows) if row[0] == 'Loja')

    assert rows[start][:5] == [
        'Loja', 'Vendas', 'Itens', 'Faturamento', '% do faturamento',
    ]
    # A loja que mais vendeu primeiro.
    assert rows[start + 1][:5] == ['Loja Norte', 2, 9, 190.0,
                                   pytest.approx(190 / 300)]
    assert rows[start + 2][:5] == ['Loja Centro', 2, 3, 110.0,
                                   pytest.approx(110 / 300)]
    assert len(sheet._charts) == 1


def test_products_ranked_by_revenue(workbook):
    rows = values(workbook['Produtos'])

    assert rows == [
        ['Produto', 'Quantidade', 'Faturamento', '% do faturamento'],
        ['Caneca', 7, 210.0, pytest.approx(0.7)],
        ['Camiseta', 1, 50.0, pytest.approx(50 / 300)],
        ['Prato', 4, 40.0, pytest.approx(40 / 300)],
    ]


def test_revenue_by_month_with_a_chart(workbook):
    sheet = workbook['Por mês']
    rows = values(sheet)

    assert rows[0] == ['Mês', 'Vendas', 'Itens', 'Faturamento']
    assert [row[0].date() for row in rows[1:]] == [
        date(2026, 9, 1), date(2026, 10, 1),
    ]
    assert [row[1:] for row in rows[1:]] == [[1, 2, 60.0], [3, 10, 240.0]]
    assert sheet['A2'].number_format == 'mm/yyyy'
    assert len(sheet._charts) == 1


def test_all_sales_together(workbook):
    sheet = workbook['Vendas']
    rows = values(sheet)

    assert rows[0] == ['Data', 'Loja', 'Produto', 'Quantidade',
                       'Preço unitário', 'Total', 'Planilha', 'Linha']
    assert rows[1][1:] == ['Loja Centro', 'Caneca', 2, 30.0, 60.0,
                           'centro.xlsx', 2]
    assert sheet['A2'].number_format == 'dd/mm/yyyy'
    assert sheet['F2'].number_format == CURRENCY
    # O cabeçalho fica parado ao rolar a lista.
    assert sheet.freeze_panes == 'A2'


def test_problems_list(workbook):
    rows = values(workbook['Problemas'])

    assert rows == [
        ['Planilha', 'Linha', 'Problema'],
        ['norte.csv', 4, 'Produto vazio'],
        ['quebrada.xlsx', None, 'arquivo ignorado: a planilha está vazia'],
    ]


def test_no_problems(tmp_path):
    path = tmp_path / 'relatorio.xlsx'
    write_report(path, SALES, [], FILES)

    rows = values(load_workbook(path)['Problemas'])

    assert rows[1] == ['Nenhum problema encontrado.', None, None]


def test_headers_are_bold(workbook):
    for name in ['Produtos', 'Vendas', 'Problemas']:
        assert workbook[name]['A1'].font.bold
