"""Gera as planilhas de exemplo em exemplos/vendas/.

Cada "loja" manda do seu jeito, como acontece de verdade: nomes de coluna
diferentes, CSV do Excel em português, datas como número de série,
título acima do cabeçalho, linhas vazias no fim e alguns erros de
digitação. Rode de novo para recriar: python exemplos/gerar_exemplos.py
"""

import csv
import random
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import Workbook

FOLDER = Path(__file__).parent / 'vendas'
PRODUCTS = {
    'Caneca personalizada': 39.90,
    'Camiseta estampada': 59.90,
    'Agenda 2027': 34.50,
    'Caderno universitário': 24.90,
    'Garrafa térmica': 79.90,
    'Ecobag': 19.90,
    'Chaveiro': 9.90,
    'Quadro decorativo': 89.90,
}
START = date(2026, 8, 1)
DAYS = 92  # agosto, setembro e outubro


def random_sales(rng, count):
    names = list(PRODUCTS)
    weights = [9, 7, 5, 6, 3, 8, 10, 2]
    for _ in range(count):
        product = rng.choices(names, weights)[0]
        day = START + timedelta(days=rng.randrange(DAYS))
        yield day, product, rng.choice([1, 1, 1, 2, 2, 3, 5]), \
            PRODUCTS[product]


def brl(value):
    return f'{value:,.2f}'.replace(',', '_').replace('.', ',').replace(
        '_', '.')


def centro(rng):
    # Excel caprichado: título, data e preço como valores de verdade.
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Vendas'
    sheet.append(['Relatório de vendas - Loja Centro'])
    sheet.append(['Exportado em 01/11/2026'])
    sheet.append([])
    sheet.append(['Data', 'Produto', 'Qtd', 'Preço unitário'])
    for day, product, quantity, price in sorted(random_sales(rng, 180)):
        sheet.append([datetime(day.year, day.month, day.day), product,
                      quantity, price])
    workbook.save(FOLDER / 'loja_centro.xlsx')


def shopping(rng):
    # CSV do Excel em português: ponto e vírgula, cp1252, "1.234,56" e
    # nomes de coluna diferentes. Com total, e alguns erros.
    rows = [['Dt. venda', 'Descrição do produto', 'Quantidade',
             'Valor unitário', 'Valor total']]
    for day, product, quantity, price in sorted(random_sales(rng, 240)):
        if rng.random() < 0.3:
            product = product.lower()
        rows.append([f'{day:%d/%m/%y}', product, str(quantity), brl(price),
                     brl(quantity * price)])
    rows[15][4] = brl(999)            # total digitado errado
    rows[42][1] = ''                  # produto em branco
    rows[77][2] = 'dois'              # quantidade por extenso
    rows[120][0] = '31/09/26'         # setembro não tem dia 31
    with open(FOLDER / 'loja_shopping.csv', 'w', newline='',
              encoding='cp1252') as file:
        csv.writer(file, delimiter=';').writerows(rows)


def bairro(rng):
    # Planilha feita à mão: só o total, datas como número de série do Excel
    # e linhas vazias no fim.
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(['DATA DA VENDA', 'ITEM', 'QTDE', 'TOTAL'])
    for day, product, quantity, price in sorted(random_sales(rng, 110)):
        serial = (day - date(1899, 12, 30)).days
        sheet.append([serial, product.upper(), quantity,
                      round(quantity * price, 2)])
    sheet.append([serial, 'ECOBAG', -1, 19.90])     # quantidade negativa
    sheet.append([None, None, None, None])
    sheet.append(['', '', '', ''])
    workbook.save(FOLDER / 'vendas_bairro.xlsx')


def unrelated():
    # Uma planilha que não é de vendas, esquecida na mesma pasta.
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(['Produto', 'Estoque mínimo'])
    sheet.append(['Caneca personalizada', 20])
    workbook.save(FOLDER / 'estoque.xlsx')


def main():
    FOLDER.mkdir(exist_ok=True)
    rng = random.Random(2026)  # sempre as mesmas planilhas
    centro(rng)
    shopping(rng)
    bairro(rng)
    unrelated()
    print(f'Planilhas criadas em {FOLDER}')


if __name__ == '__main__':
    main()
